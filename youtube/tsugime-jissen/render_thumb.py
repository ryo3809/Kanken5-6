# -*- coding: utf-8 -*-
"""thumb.json -> thumb.pptx （1枚・16:9・YouTubeサムネイル／右にアバター）

thumb.json の "avatar" に画像パスを入れると、右パネルにカバークロップで配置する。
パスが無い／ファイルが見つからない場合は、同じ位置にプレースホルダ枠を描く。
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

W, H = Inches(13.333), Inches(7.5)
BG    = RGBColor(0x1C, 0x1A, 0x19)
PANEL = RGBColor(0x2A, 0x26, 0x24)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY  = RGBColor(0xBD, 0xB6, 0xAF)
FONT  = "Hiragino Sans"

# 左：テキスト、右：アバター
TEXT_X, TEXT_W = Inches(1.15), Inches(6.45)
AV_X, AV_Y, AV_W, AV_H = Inches(8.05), Inches(0.6), Inches(4.75), Inches(6.3)


def ea(run, name=FONT):
    rPr = run.font._rPr
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def box(s, x, y, w, h, text, size, color, bold=True, align=PP_ALIGN.LEFT,
        spacing=1.2, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = spacing
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    ea(r)
    return tb


def rect(s, x, y, w, h, fill, radius=None, line=None, dash=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius is not None:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(2)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    return shp


def place_avatar(s, path, accent):
    """枠いっぱいにカバークロップして配置。無ければプレースホルダ枠。"""
    rect(s, AV_X, AV_Y, AV_W, AV_H, PANEL, radius=0.06)
    if path and os.path.exists(path):
        from PIL import Image
        iw, ih = Image.open(path).size
        frame_ratio, img_ratio = AV_W / AV_H, iw / ih
        pic = s.shapes.add_picture(path, AV_X, AV_Y, AV_W, AV_H)
        if img_ratio > frame_ratio:            # 画像が横長 → 左右を削る
            keep = frame_ratio / img_ratio
            pic.crop_left = pic.crop_right = (1 - keep) / 2
        elif img_ratio < frame_ratio:          # 画像が縦長 → 下を多めに削って顔を上に残す
            keep = img_ratio / frame_ratio
            pic.crop_top = (1 - keep) * 0.25
            pic.crop_bottom = (1 - keep) * 0.75
        rect(s, AV_X, AV_Y, AV_W, Inches(0.09), accent)   # 画像の上に重ねる
        return True
    rect(s, AV_X, AV_Y, AV_W, Inches(0.09), accent)
    rect(s, AV_X + Inches(0.3), AV_Y + Inches(0.3), AV_W - Inches(0.6), AV_H - Inches(0.6),
         None, radius=0.05, line=RGBColor(0x6E, 0x67, 0x60), dash=True)
    tb = box(s, AV_X, AV_Y + AV_H / 2 - Inches(0.55), AV_W, Inches(1.1),
             "アバター画像", 20, GREY, bold=False, align=PP_ALIGN.CENTER, spacing=1.5)
    p2 = tb.text_frame.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = "thumb.json の avatar にパスを指定"
    r2.font.size = Pt(14); r2.font.color.rgb = GREY; r2.font.name = FONT
    ea(r2)
    return False


def main(cfg_path, out_path):
    c = json.load(open(cfg_path, encoding="utf-8"))
    accent = RGBColor.from_string(c.get("accent", "D65C4C"))
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, BG)
    rect(s, 0, 0, Inches(0.34), H, accent)

    if c.get("badge"):
        bw = Inches(1.75)
        rect(s, TEXT_X, Inches(0.95), bw, Inches(0.6), accent, radius=0.4)
        box(s, TEXT_X, Inches(1.09), bw, Inches(0.4), c["badge"], 21, WHITE, align=PP_ALIGN.CENTER)

    # 2行組み：上段は白、下段は強調色（1行8文字までが基本）
    def size_for(t): return 62 if len(t) <= 6 else 54 if len(t) <= 8 else 46 if len(t) <= 10 else 40
    sz = min(size_for(c["line1"]), size_for(c["line2"]))
    box(s, TEXT_X, Inches(2.05), TEXT_W, Inches(1.3), c["line1"], sz, WHITE)
    box(s, TEXT_X, Inches(3.35), TEXT_W, Inches(1.3), c["line2"], sz, accent)

    rect(s, TEXT_X + Inches(0.03), Inches(4.95), Inches(2.1), Inches(0.09), accent)
    box(s, TEXT_X, Inches(5.25), TEXT_W, Inches(1.0), c.get("sub", ""), 20, GREY,
        bold=False, spacing=1.35)
    box(s, TEXT_X, Inches(6.4), TEXT_W, Inches(0.5), c.get("channel", ""), 19, GREY, bold=False)

    placed = place_avatar(s, c.get("avatar"), accent)
    prs.save(out_path)
    print("wrote", out_path, "/ avatar:", "placed" if placed else "PLACEHOLDER（画像未指定）")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
