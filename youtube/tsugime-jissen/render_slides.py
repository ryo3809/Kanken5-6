# -*- coding: utf-8 -*-
"""deck.json -> slides.pptx （ロジた標準レイアウト：title/profile/compare/statement/flow/rows/spotlight/closing）"""
import json, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

W, H = Inches(13.333), Inches(7.5)
BG      = RGBColor(0xF7, 0xF5, 0xF2)
INK     = RGBColor(0x22, 0x20, 0x1E)
MUTED   = RGBColor(0x6E, 0x67, 0x60)
LINE    = RGBColor(0xDD, 0xD7, 0xD0)
CARD    = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x26, 0x23, 0x21)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

def rgb(hexstr): return RGBColor.from_string(hexstr)

class Deck:
    def __init__(self, meta):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.accent = rgb(meta.get("accent", "B85042"))
        self.font = meta.get("font", "Hiragino Sans")
        self.channel = meta.get("channel", "")

    # ---------- primitives ----------
    def _slide(self, dark=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
        bg.fill.solid(); bg.fill.fore_color.rgb = DARK if dark else BG
        bg.line.fill.background(); bg.shadow.inherit = False
        return s

    @staticmethod
    def _ea_font_name(rPr, name):
        from pptx.oxml.ns import qn
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = rPr.makeelement(qn(tag), {})
                rPr.append(el)
            el.set("typeface", name)

    def _ea_font(self, run):
        self._ea_font_name(run.font._rPr, self.font)

    def _box(self, s, x, y, w, h, text, size=18, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=Pt(6), line=1.28):
        tb = s.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for i, ln in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = space
            p.line_spacing = line
            r = p.add_run(); r.text = ln
            f = r.font
            f.size = Pt(size); f.bold = bold; f.color.rgb = color
            f.name = self.font
            self._ea_font(r)
        return tb

    def _rect(self, s, x, y, w, h, fill=CARD, line=None, width=Pt(1), radius=None):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
        if radius is not None:
            shp.adjustments[0] = radius
        if fill is None: shp.fill.background()
        else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if line is None: shp.line.fill.background()
        else: shp.line.color.rgb = line; shp.line.width = width
        shp.shadow.inherit = False
        return shp

    def _accent_bar(self, s, x, y, w=Inches(0.62), h=Inches(0.08)):
        self._rect(s, x, y, w, h, fill=self.accent)

    def _heading(self, s, text, dark=False):
        self._accent_bar(s, Inches(0.9), Inches(0.72))
        self._box(s, Inches(0.9), Inches(0.95), Inches(11.6), Inches(0.9), text,
                  size=32, bold=True, color=WHITE if dark else INK)
        self._rect(s, Inches(0.9), Inches(1.78), Inches(11.55), Emu(9525), fill=LINE)

    def _footer(self, s, n, dark=False):
        self._box(s, Inches(0.9), Inches(6.83), Inches(6.0), Inches(0.3), self.channel,
                  size=11, color=MUTED)
        self._box(s, Inches(11.4), Inches(6.83), Inches(1.05), Inches(0.3), f"{n:02d}",
                  size=11, color=MUTED, align=PP_ALIGN.RIGHT)

    # ---------- layouts ----------
    def title(self, d, n):
        s = self._slide()
        self._rect(s, 0, 0, Inches(0.28), H, fill=self.accent)
        self._box(s, Inches(1.15), Inches(1.55), Inches(11.0), Inches(0.4),
                  d.get("kicker", ""), size=16, bold=True, color=self.accent)
        self._box(s, Inches(1.15), Inches(2.25), Inches(11.2), Inches(2.6),
                  d["title"], size=44, bold=True, color=INK, line=1.35)
        self._rect(s, Inches(1.15), Inches(4.92), Inches(2.0), Inches(0.05), fill=self.accent)
        self._box(s, Inches(1.15), Inches(5.25), Inches(11.0), Inches(0.6),
                  d.get("subtitle", ""), size=20, color=MUTED)
        self._footer(s, n)

    def profile(self, d, n):
        s = self._slide()
        self._heading(s, d.get("heading", "自己紹介"))
        self._rect(s, Inches(0.9), Inches(2.35), Inches(3.15), Inches(3.5), fill=CARD, line=LINE)
        self._rect(s, Inches(0.9), Inches(2.35), Inches(3.15), Inches(0.09), fill=self.accent)
        self._box(s, Inches(1.15), Inches(3.5), Inches(2.65), Inches(1.0), d.get("name", ""),
                  size=34, bold=True, color=INK, align=PP_ALIGN.CENTER)
        y = Inches(2.55)
        for ln in d["lines"]:
            self._rect(s, Inches(4.5), y + Inches(0.16), Inches(0.11), Inches(0.11), fill=self.accent)
            self._box(s, Inches(4.85), y, Inches(7.6), Inches(1.0), ln, size=19, color=INK)
            y += Inches(1.12)
        self._footer(s, n)

    def compare(self, d, n):
        s = self._slide()
        self._heading(s, d["heading"])
        cw, gap = Inches(5.5), Inches(0.55)
        x0 = Inches(0.9); x1 = x0 + cw + gap
        top = Inches(2.15); ch = Inches(3.72)
        for x, col, is_right in ((x0, d["left"], False), (x1, d["right"], True)):
            self._rect(s, x, top, cw, ch, fill=CARD, line=LINE)
            self._rect(s, x, top, cw, Inches(0.09), fill=self.accent if is_right else RGBColor(0xB9,0xB2,0xAA))
            mark_c = self.accent if is_right else MUTED
            self._box(s, x + Inches(0.35), top + Inches(0.34), Inches(0.7), Inches(0.6),
                      col.get("mark", ""), size=26, bold=True, color=mark_c)
            self._box(s, x + Inches(1.05), top + Inches(0.42), cw - Inches(1.4), Inches(0.5),
                      col.get("label", ""), size=18, bold=True, color=mark_c)
            n_items = len(col["items"])
            y = top + Inches(1.25) + (Inches(0.34) if n_items <= 2 else Inches(0))
            for it in col["items"]:
                self._rect(s, x + Inches(0.38), y + Inches(0.14), Inches(0.1), Inches(0.1), fill=LINE)
                self._box(s, x + Inches(0.68), y, cw - Inches(1.1), Inches(0.9), it, size=17, color=INK)
                y += Inches(0.82)
        if d.get("footnote"):
            self._box(s, Inches(0.9), Inches(6.12), Inches(11.5), Inches(0.55), d["footnote"],
                      size=16, bold=True, color=self.accent)
        self._footer(s, n)

    def statement(self, d, n):
        s = self._slide(dark=True)
        self._rect(s, Inches(0.9), Inches(2.15), Inches(0.09), Inches(2.6), fill=self.accent)
        self._box(s, Inches(1.35), Inches(2.15), Inches(11.0), Inches(2.6), d["text"],
                  size=40, bold=True, color=WHITE, line=1.4)
        if d.get("note"):
            self._box(s, Inches(1.35), Inches(5.05), Inches(11.0), Inches(0.7), d["note"],
                      size=18, color=RGBColor(0xC9,0xC2,0xBA))
        self._box(s, Inches(0.9), Inches(6.83), Inches(6.0), Inches(0.3), self.channel,
                  size=11, color=RGBColor(0x8C,0x85,0x7E))
        self._box(s, Inches(11.4), Inches(6.83), Inches(1.05), Inches(0.3), f"{n:02d}",
                  size=11, color=RGBColor(0x8C,0x85,0x7E), align=PP_ALIGN.RIGHT)

    def rows(self, d, n):
        s = self._slide()
        self._heading(s, d["heading"])
        items = d["rows"]
        top = Inches(2.12)
        avail = Inches(4.02) if d.get("source") else Inches(4.6)
        gap = Inches(0.14)
        rh = int((avail - gap * (len(items) - 1)) / len(items))
        cap = Inches(1.45)
        if rh > cap:                      # 行が少ないときは伸ばしすぎず、縦中央に寄せる
            rh = cap
            used = rh * len(items) + gap * (len(items) - 1)
            top = top + int((avail - used) / 2)
        dense = len(items) >= 4
        t_size, d_size = (18, 13) if dense else (21, 15)
        t_off, d_off = (Inches(0.13), Inches(0.47)) if dense else (Inches(0.2), Inches(0.72))
        y = top
        for it in items:
            self._rect(s, Inches(0.9), y, Inches(11.55), rh, fill=CARD, line=LINE)
            self._rect(s, Inches(0.9), y, Inches(0.075), rh, fill=self.accent)
            self._box(s, Inches(1.25), y + t_off + Inches(0.02), Inches(0.85), Inches(0.45),
                      it["n"], size=t_size - 2, bold=True, color=self.accent)
            self._box(s, Inches(2.1), y + t_off, Inches(10.15), Inches(0.45),
                      it["title"], size=t_size, bold=True, color=INK)
            if it.get("desc"):
                self._box(s, Inches(2.1), y + d_off, Inches(10.15), Inches(0.4),
                          it["desc"], size=d_size, color=MUTED)
            y += rh + gap
        if d.get("source"):
            self._box(s, Inches(0.9), Inches(6.34), Inches(11.5), Inches(0.35), "出典：" + d["source"],
                      size=12, color=MUTED)
        self._footer(s, n)

    def flow(self, d, n):
        s = self._slide()
        self._heading(s, d["heading"])
        steps = d["steps"]; joints = d.get("joints", [])
        cw = Inches(3.62); gap = Inches(0.72)
        x = Inches(0.9); top = Inches(2.55); ch = Inches(2.55)
        chips = []
        for i, st in enumerate(steps):
            dark = bool(st.get("dark"))
            self._rect(s, x, top, cw, ch, fill=DARK if dark else CARD, line=None if dark else LINE)
            self._rect(s, x, top, cw, Inches(0.09), fill=self.accent)
            self._box(s, x + Inches(0.32), top + Inches(0.42), cw - Inches(0.64), Inches(0.75),
                      st["title"], size=22, bold=True, color=WHITE if dark else INK)
            self._box(s, x + Inches(0.32), top + Inches(1.32), cw - Inches(0.64), Inches(1.05),
                      st.get("desc", ""), size=15,
                      color=RGBColor(0xD8,0xD2,0xCC) if dark else MUTED)
            if i < len(steps) - 1:
                ax = x + cw + Inches(0.12)
                ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, top + Inches(1.08), Inches(0.48), Inches(0.34))
                ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xC6,0xBE,0xB6)
                ar.line.fill.background(); ar.shadow.inherit = False
                if i < len(joints):
                    chips.append((ax + Inches(0.24), joints[i]))
            x += cw + gap
        for cx, label in chips:
            cwid = Inches(1.62)
            cy = top + ch + Inches(0.16)
            self._rect(s, cx - cwid / 2, cy, cwid, Inches(0.4), fill=self.accent, radius=0.5)
            self._box(s, cx - cwid / 2, cy + Inches(0.1), cwid, Inches(0.3),
                      label, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        if d.get("footnote"):
            fy = Inches(5.95) if joints else Inches(5.72)
            self._box(s, Inches(0.9), fy, Inches(11.5), Inches(0.6), d["footnote"],
                      size=17, bold=True, color=self.accent)
        self._footer(s, n)

    def spotlight(self, d, n):
        s = self._slide()
        self._accent_bar(s, Inches(0.9), Inches(0.72))
        self._box(s, Inches(0.9), Inches(0.95), Inches(4.0), Inches(0.5), d["index"],
                  size=20, bold=True, color=self.accent)
        self._box(s, Inches(0.9), Inches(1.55), Inches(11.5), Inches(1.15), d["title"],
                  size=40, bold=True, color=INK)
        self._rect(s, Inches(0.9), Inches(2.92), Inches(11.55), Emu(9525), fill=LINE)
        y = Inches(3.3)
        for b in d["body"]:
            self._rect(s, Inches(0.95), y + Inches(0.2), Inches(0.22), Inches(0.06), fill=self.accent)
            self._box(s, Inches(1.45), y, Inches(10.9), Inches(0.8), b, size=21, color=INK)
            y += Inches(0.78)
        self._footer(s, n)

    def closing(self, d, n):
        s = self._slide()
        self._heading(s, d["heading"])
        y = Inches(2.3)
        for i, ln in enumerate(d["lines"]):
            self._box(s, Inches(0.95), y + Inches(0.02), Inches(0.55), Inches(0.5), f"{i+1}.",
                      size=19, bold=True, color=self.accent)
            self._box(s, Inches(1.6), y, Inches(10.8), Inches(0.8), ln, size=20, color=INK)
            y += Inches(0.86)
        self._footer(s, n)

    def build(self, slides):
        for i, sl in enumerate(slides, 1):
            getattr(self, sl["layout"])(sl, i)
        return self.prs


def main(deck_path, out_path):
    d = json.load(open(deck_path, encoding="utf-8"))
    deck = Deck(d.get("meta", {}))
    deck.build(d["slides"]).save(out_path)
    print("wrote", out_path, len(d["slides"]), "slides")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
